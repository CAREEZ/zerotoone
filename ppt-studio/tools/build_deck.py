#!/usr/bin/env python3
"""spec(JSON) → 편집 가능한 .pptx 생성기.

사용법:
    python3 ppt-studio/tools/build_deck.py <spec.json> <out.pptx>

- 모든 텍스트는 텍스트 상자, 차트는 네이티브 PowerPoint 차트, 표는 네이티브 표로 생성한다.
  (PowerPoint에서 바로 수정 가능 / 차트는 '데이터 편집'으로 수치 수정 가능)
- 브랜드 토큰은 ppt-studio/design-system/brands/<brand>.json 에서 읽는다.
- 텍스트에 "[확인필요]"가 들어 있으면 경고색으로 표시하고, 빌드 후 목록을 출력한다.
- 분량 규칙(제목 길이, 불릿 개수 등)을 어기면 경고를 출력한다. 빌드는 계속한다.
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
BRAND_DIR = ROOT / "design-system" / "brands"
CHECK_TAG = "[확인필요]"

SIZES = {
    "16:9": (Inches(13.333), Inches(7.5)),
    "A4-landscape": (Emu(10692000), Emu(7560000)),  # 297 x 210 mm
    "A4-portrait": (Emu(7560000), Emu(10692000)),   # 210 x 297 mm
}

LIMITS = {"title": 40, "bullets": 4, "bullet_len": 40, "kpi": 4, "steps": 5, "table_rows": 8}
DARK_TYPES = {"cover", "closing"}  # 어두운 배경은 표지·마무리만

warnings, checks = [], []


def split_keyword(line):
    """'키워드: 내용' → ('키워드', '내용'). 형식이 아니면 (None, line)."""
    s = str(line)
    if ": " in s:
        k, v = s.split(": ", 1)
        if 0 < len(k) <= 15:
            return k, v
    return None, s


def lint_sequence(spec):
    """슬라이드 순서 규칙 점검 — 같은 레이아웃 연속 사용 금지 (간지 제외)."""
    prev = None
    for i, d in enumerate(spec["slides"], start=1):
        t = d["type"]
        if t == prev and t not in ("section",):
            warnings.append(f"p{i} '{t}' 레이아웃이 직전 장과 같음 — 연속 사용 금지")
        prev = t


def rgb(hex_str):
    return RGBColor.from_string(hex_str.lstrip("#"))


def load_brand(name):
    path = BRAND_DIR / f"{name}.json"
    if not path.exists():
        sys.exit(f"브랜드 파일 없음: {path}")
    brand = json.loads(path.read_text(encoding="utf-8"))
    if brand.get("status", "").startswith("미확인"):
        warnings.append(f"브랜드 '{name}' 색상·로고가 미확인 상태임 ({brand['status']})")
    return brand


class Deck:
    def __init__(self, spec):
        self.spec = spec
        self.brand = load_brand(spec.get("brand", "default"))
        self.c = self.brand["colors"]
        self.font = self.brand["fonts"]["body"]
        self.font_head = self.brand["fonts"]["heading"]
        self.prs = Presentation()
        w, h = SIZES[spec.get("size", "16:9")]
        self.prs.slide_width, self.prs.slide_height = w, h
        self.W, self.H = w, h
        self.portrait = h > w
        self.M = int(w * 0.045)  # 좌우 여백
        self.page = 0

    # ---------- 기본 도형 ----------
    def _set_font(self, run, size, color, bold=False, heading=False):
        name = self.font_head if heading else self.font
        f = run.font
        f.name, f.size, f.bold = name, Pt(size), bold
        f.color.rgb = rgb(color)
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):  # 한글(동아시아) 글꼴까지 지정
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set("typeface", name)

    def text(self, slide, x, y, w, h, lines, size=16, color=None, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, heading=False, spacing=1.15,
             bullet=False, where=""):
        color = color or self.c["text"]
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        if isinstance(lines, str):
            lines = [lines]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            if bullet:
                p.space_after = Pt(size * 0.6)
            col = color
            if CHECK_TAG in str(line):
                col = self.c["warning"]
                checks.append(f"p{self.page} {where}: {line}")
            key, rest = split_keyword(line) if bullet else (None, str(line))
            if key:  # '키워드: 내용' → 키워드만 굵게
                r1 = p.add_run()
                r1.text = f"•  {key}  "
                self._set_font(r1, size, self.c["primary"] if col == color else col, True, heading)
            run = p.add_run()
            run.text = ("" if key else ("•  " if bullet else "")) + rest
            self._set_font(run, size, col, bold, heading)
        return box

    def rect(self, slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, x, y, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def new_slide(self, bg=None):
        self.page += 1
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(slide, 0, 0, self.W, self.H, bg or self.c["background"])
        return slide

    # ---------- 공통 프레임: 챕터 / 제목 / 푸터 (모든 본문 슬라이드 동일 위치) ----------
    def frame(self, s, d):
        M, W, H = self.M, self.W, self.H
        title = d.get("title", "")
        if len(title) > LIMITS["title"]:
            warnings.append(f"p{self.page} 제목 {len(title)}자 > {LIMITS['title']}자: {title}")
        if d.get("chapter"):
            self.text(s, M, int(H * 0.055), W - 2 * M, int(H * 0.04), d["chapter"],
                      size=11, color=self.c["accent"], bold=True, where="chapter")
        self.text(s, M, int(H * 0.095), W - 2 * M, int(H * 0.11), title,
                  size=26 if not self.portrait else 22, bold=True, heading=True,
                  color=self.c["primary"], anchor=MSO_ANCHOR.TOP, spacing=1.1, where="title")
        if d.get("subtitle"):
            self.text(s, M, int(H * 0.2), W - 2 * M, int(H * 0.05), d["subtitle"],
                      size=14, color=self.c["muted"], where="subtitle")
        # 푸터: 발행 법인 · 출처 · 페이지
        fy = H - int(H * 0.06)
        self.rect(s, M, fy - Emu(60000), W - 2 * M, Emu(9525), self.c["line"])
        foot = "   |   ".join(x for x in [f"출처: {d['source']}" if d.get("source") else "", self.brand["footer"]] if x)
        self.text(s, M, fy, W - 2 * M - Inches(1), int(H * 0.035), foot,
                  size=9, color=self.c["muted"], where="footer")
        self.text(s, W - M - Inches(1), fy, Inches(1), int(H * 0.035), str(self.page),
                  size=9, color=self.c["muted"], align=PP_ALIGN.RIGHT)
        if d.get("notes"):
            s.notes_slide.notes_text_frame.text = d["notes"]
        return int(H * 0.27)  # 본문 시작 y

    def body_box(self, top):
        M, W, H = self.M, self.W, self.H
        return M, top, W - 2 * M, H - top - int(H * 0.1)

    # ---------- 슬라이드 유형 ----------
    def cover(self, d):
        s = self.new_slide(self.c["primary"])
        M, W, H = self.M, self.W, self.H
        self.rect(s, M, int(H * 0.38), Inches(0.9), Emu(57150), self.c["accent"])
        self.text(s, M, int(H * 0.42), W - 2 * M, int(H * 0.22), d["title"], size=40 if not self.portrait else 32,
                  bold=True, heading=True, color="FFFFFF", spacing=1.1, where="cover")
        if d.get("subtitle"):
            self.text(s, M, int(H * 0.64), W - 2 * M, int(H * 0.08), d["subtitle"], size=18,
                      color=self.c["on_primary_muted"], where="cover")
        meta = " · ".join(x for x in [d.get("date"), d.get("presenter"), self.brand["company"]] if x)
        self.text(s, M, H - int(H * 0.12), W - 2 * M, int(H * 0.05), meta, size=12,
                  color=self.c["on_primary_muted"], where="cover")
        if d.get("notes"):
            s.notes_slide.notes_text_frame.text = d["notes"]

    def section(self, d):
        s = self.new_slide(self.c["surface"])  # 간지도 밝은 배경 (어두운 배경은 표지·마무리만)
        M, W, H = self.M, self.W, self.H
        self.text(s, M, int(H * 0.3), W - 2 * M, int(H * 0.2), d.get("no", ""), size=72, bold=True,
                  heading=True, color=self.c["accent"])
        self.rect(s, M, int(H * 0.52), Inches(0.9), Emu(57150), self.c["primary"])
        self.text(s, M, int(H * 0.56), W - 2 * M, int(H * 0.12), d["title"], size=32, bold=True,
                  heading=True, color=self.c["primary"], where="section")

    def agenda(self, d):
        s = self.new_slide()
        top = self.frame(s, {"title": d.get("title", "목차"), **{k: v for k, v in d.items() if k != "title"}})
        x, y, w, h = self.body_box(top)
        items = d["items"]
        row = min(int(h / max(len(items), 1)), Inches(0.9))
        for i, it in enumerate(items):
            yy = y + i * row
            self.text(s, x, yy, Inches(0.9), row, f"{i + 1:02d}", size=22, bold=True,
                      color=self.c["accent"], heading=True)
            self.text(s, x + Inches(1.0), yy + Emu(40000), w - Inches(1.0), row, it, size=18,
                      color=self.c["text"], where="agenda")
            self.rect(s, x, yy + row - Emu(30000), w, Emu(6350), self.c["line"])

    def bullets(self, d):
        s = self.new_slide()
        top = self.frame(s, d)
        x, y, w, h = self.body_box(top)
        items = d.get("bullets", [])
        if len(items) > LIMITS["bullets"]:
            warnings.append(f"p{self.page} 불릿 {len(items)}개 > {LIMITS['bullets']}개 — 쪼개거나 표로 전환 검토")
        for b in items:
            if len(b) > LIMITS["bullet_len"]:
                warnings.append(f"p{self.page} 불릿 {len(b)}자 > {LIMITS['bullet_len']}자: {b[:30]}…")
        if d.get("message"):
            self.rect(s, x, y, Emu(57150), Inches(0.55), self.c["accent"])
            self.text(s, x + Inches(0.25), y, w - Inches(0.25), Inches(0.6), d["message"], size=20,
                      bold=True, color=self.c["text"], anchor=MSO_ANCHOR.MIDDLE, where="message")
            y += Inches(1.0)
        self.text(s, x, y, w, h - (y - top), items, size=18, color=self.c["text"], bullet=True,
                  spacing=1.25, where="bullets")

    def kpi(self, d):
        s = self.new_slide()
        top = self.frame(s, d)
        x, y, w, h = self.body_box(top)
        items = d["items"]
        if len(items) > LIMITS["kpi"]:
            warnings.append(f"p{self.page} KPI {len(items)}개 > {LIMITS['kpi']}개")
        n = len(items)
        cols = n if not self.portrait else min(n, 2)
        rows = (n + cols - 1) // cols
        gap = Inches(0.3)
        cw = int((w - gap * (cols - 1)) / cols)
        ch = min(int((h - gap * (rows - 1)) / rows), Inches(3.2))
        for i, it in enumerate(items):
            cx = x + (i % cols) * (cw + gap)
            cy = y + (i // cols) * (ch + gap)
            self.rect(s, cx, cy, cw, ch, self.c["surface"])
            self.rect(s, cx, cy, cw, Emu(57150), self.c["accent"])
            pad = Inches(0.3)
            self.text(s, cx + pad, cy + pad, cw - 2 * pad, int(ch * 0.4), it["value"], size=40,
                      bold=True, heading=True, color=self.c["primary"], where="kpi")
            self.text(s, cx + pad, cy + int(ch * 0.5), cw - 2 * pad, int(ch * 0.15), it["label"], size=16,
                      bold=True, color=self.c["text"], where="kpi")
            if it.get("desc"):
                self.text(s, cx + pad, cy + int(ch * 0.66), cw - 2 * pad, int(ch * 0.3), it["desc"], size=12,
                          color=self.c["muted"], where="kpi")

    def chart(self, d):
        s = self.new_slide()
        top = self.frame(s, d)
        x, y, w, h = self.body_box(top)
        kind = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "stacked": XL_CHART_TYPE.COLUMN_STACKED,
            "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.DOUGHNUT,
        }[d.get("chart", "bar")]
        cd = CategoryChartData()
        cd.categories = d["categories"]
        for ser in d["series"]:
            cd.add_series(ser["name"], ser["values"])
        side = d.get("takeaway") and not self.portrait
        cw = int(w * 0.66) if side else w
        chh = h if not (d.get("takeaway") and self.portrait) else int(h * 0.7)
        gf = s.shapes.add_chart(kind, x, y, cw, chh, cd)
        ch = gf.chart
        ch.has_title = False
        ch.font.size, ch.font.name = Pt(12), self.font
        ch.font.color.rgb = rgb(self.c["text"])
        palette = self.c["chart"]
        multi = len(d["series"]) > 1 or kind == XL_CHART_TYPE.DOUGHNUT
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        plot = ch.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(11)
        dl.number_format = d.get("number_format", "#,##0")
        dl.number_format_is_linked = False
        if kind == XL_CHART_TYPE.DOUGHNUT:
            dl.number_format = "0%"
            dl.show_percentage, dl.show_value = True, False
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(palette[i % len(palette)])
        else:
            if kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED):
                plot.gap_width = 60
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            hl = d.get("highlight")  # 강조할 카테고리 인덱스 (단일 시리즈)
            for i, ser in enumerate(plot.series):
                col = palette[i % len(palette)]
                if kind == XL_CHART_TYPE.LINE_MARKERS:
                    ser.format.line.color.rgb = rgb(col)
                    ser.format.line.width = Pt(2.5)
                    ser.smooth = False
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = rgb(col)
                    if hl is not None and len(d["series"]) == 1:
                        for j, pt in enumerate(ser.points):
                            pt.format.fill.solid()
                            pt.format.fill.fore_color.rgb = rgb(self.c["accent"] if j == hl else self.c["chart_muted"])
            va = ch.value_axis
            va.has_major_gridlines = False
            va.visible = False
            ca = ch.category_axis
            ca.format.line.color.rgb = rgb(self.c["line"])
            ca.tick_labels.font.size = Pt(12)
        if d.get("unit"):
            self.text(s, x, y - Inches(0.35), cw, Inches(0.3), f"단위: {d['unit']}", size=10,
                      color=self.c["muted"], where="unit")
        if d.get("takeaway"):
            if side:
                tx, ty, tw, th = x + cw + Inches(0.4), y + Inches(0.3), w - cw - Inches(0.4), h - Inches(0.3)
            else:
                tx, ty, tw, th = x, y + chh + Inches(0.2), w, h - chh - Inches(0.2)
            self.rect(s, tx, ty, Emu(57150), Inches(0.5), self.c["accent"])
            self.text(s, tx + Inches(0.25), ty, tw - Inches(0.25), th, d["takeaway"], size=16,
                      color=self.c["text"], spacing=1.3, where="takeaway")

    def table(self, d):
        s = self.new_slide()
        top = self.frame(s, d)
        x, y, w, h = self.body_box(top)
        cols, rows = d["columns"], d["rows"]
        if len(rows) > LIMITS["table_rows"]:
            warnings.append(f"p{self.page} 표 {len(rows)}행 > {LIMITS['table_rows']}행 — 부록 분리 검토")
        rh = Inches(0.5)
        shape = s.shapes.add_table(len(rows) + 1, len(cols), x, y, w, rh * (len(rows) + 1))
        tbl = shape.table
        tblPr = tbl._tbl.tblPr
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
        hl = d.get("highlight_col")
        for ci, name in enumerate(cols):
            self._cell(tbl.cell(0, ci), name, self.c["primary"], "FFFFFF", True, ci > 0)
        for ri, r in enumerate(rows, start=1):
            for ci, v in enumerate(r):
                fill = self.c["highlight_bg"] if ci == hl else ("FFFFFF" if ri % 2 else self.c["surface"])
                self._cell(tbl.cell(ri, ci), v, fill, self.c["text"], ci == hl or ci == 0, ci > 0)

    def _cell(self, cell, value, fill, color, bold, right):
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT if right and _is_num(value) else (PP_ALIGN.CENTER if right else PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = str(value)
        col = self.c["warning"] if CHECK_TAG in str(value) else color
        if CHECK_TAG in str(value):
            checks.append(f"p{self.page} table: {value}")
        self._set_font(run, 13, col, bold)

    def compare(self, d):
        s = self.new_slide()
        top = self.frame(s, d)
        x, y, w, h = self.body_box(top)
        gap = Inches(0.4)
        if self.portrait:
            cw, ch = w, int((h - gap) / 2)
            pos = [(x, y), (x, y + ch + gap)]
        else:
            cw, ch = int((w - gap) / 2), h
            pos = [(x, y), (x + cw + gap, y)]
        for (cx, cy), key, head in zip(pos, ("left", "right"), (self.c["muted"], self.c["primary"])):
            col = d[key]
            self.rect(s, cx, cy, cw, Inches(0.6), head)
            self.text(s, cx + Inches(0.25), cy, cw - Inches(0.5), Inches(0.6), col["title"], size=18, bold=True,
                      color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE, where="compare")
            self.rect(s, cx, cy + Inches(0.6), cw, ch - Inches(0.6), self.c["surface"])
            self.text(s, cx + Inches(0.3), cy + Inches(0.9), cw - Inches(0.6), ch - Inches(1.1), col["bullets"],
                      size=16, bullet=True, where="compare")

    def process(self, d):
        s = self.new_slide()
        top = self.frame(s, d)
        x, y, w, h = self.body_box(top)
        steps = d["steps"]
        if len(steps) > LIMITS["steps"]:
            warnings.append(f"p{self.page} 단계 {len(steps)}개 > {LIMITS['steps']}개")
        n = len(steps)
        if self.portrait:
            rh = int(h / n)
            for i, st in enumerate(steps):
                cy = y + i * rh
                self.rect(s, x, cy, Inches(0.7), Inches(0.7), self.c["primary"], MSO_SHAPE.OVAL)
                self.text(s, x, cy, Inches(0.7), Inches(0.7), str(i + 1), size=18, bold=True, color="FFFFFF",
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                self.text(s, x + Inches(1), cy, w - Inches(1), Inches(0.4), st["title"], size=16, bold=True,
                          where="process")
                self.text(s, x + Inches(1), cy + Inches(0.42), w - Inches(1), rh - Inches(0.5),
                          st.get("desc", ""), size=12, color=self.c["muted"], where="process")
            return
        gap = Inches(0.2)
        cw = int((w - gap * (n - 1)) / n)
        y += int(h * 0.15)
        for i, st in enumerate(steps):
            cx = x + i * (cw + gap)
            fill = self.c["primary"] if i == n - 1 else self.c["surface"]
            fg = "FFFFFF" if i == n - 1 else self.c["text"]
            self.rect(s, cx, y, cw, Inches(1.1), fill, MSO_SHAPE.CHEVRON if i < n - 1 else MSO_SHAPE.PENTAGON)
            self.text(s, cx + Inches(0.35), y, cw - Inches(0.6), Inches(1.1), f"{i + 1}. {st['title']}", size=15,
                      bold=True, color=fg, anchor=MSO_ANCHOR.MIDDLE, where="process")
            self.text(s, cx + Inches(0.1), y + Inches(1.35), cw - Inches(0.2), h - Inches(1.4), st.get("desc", ""),
                      size=13, color=self.c["text"], spacing=1.3, where="process")

    def closing(self, d):
        s = self.new_slide(self.c["primary"])
        M, W, H = self.M, self.W, self.H
        self.text(s, M, int(H * 0.38), W - 2 * M, int(H * 0.15), d.get("title", "감사합니다"), size=36,
                  bold=True, heading=True, color="FFFFFF", where="closing")
        if d.get("contact"):
            self.text(s, M, int(H * 0.56), W - 2 * M, int(H * 0.2), d["contact"], size=14,
                      color=self.c["on_primary_muted"], spacing=1.4, where="closing")

    def build(self, out):
        lint_sequence(self.spec)
        for d in self.spec["slides"]:
            getattr(self, d["type"])(d)
        self.prs.save(out)


def _is_num(v):
    t = str(v).replace(",", "").replace("%", "").replace("억", "").replace("원", "").replace("+", "").replace("-", "").strip()
    try:
        float(t)
        return True
    except ValueError:
        return False


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    deck = Deck(spec)
    deck.build(sys.argv[2])
    print(f"생성 완료: {sys.argv[2]} ({deck.page}장, 브랜드={spec.get('brand', 'default')}, 크기={spec.get('size', '16:9')})")
    if warnings:
        print("\n[분량·규칙 경고]")
        for w in warnings:
            print(" -", w)
    if checks:
        print(f"\n[{CHECK_TAG} 항목 {len(checks)}건 — 발표 전 반드시 확인]")
        for c in checks:
            print(" -", c)


if __name__ == "__main__":
    main()
